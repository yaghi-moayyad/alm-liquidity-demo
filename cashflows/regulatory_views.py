import csv
from decimal import Decimal
from django.db import transaction
from django.http import HttpResponse
from rest_framework.decorators import api_view
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from .models import Entity, RegulatoryConfiguration, RegulatoryMapping, RegulatoryCalculation, RegulatoryContribution
from .regulatory_engine import calculate_regulatory, REGULATORY_ENGINE_VERSION
from .regulatory_serializers import RegulatoryConfigurationSerializer, RegulatoryMappingSerializer, RegulatoryCalculationSerializer, RegulatoryContributionSerializer


def _entity(slug):
    try: return Entity.objects.get(slug=slug)
    except Entity.DoesNotExist: raise ValidationError({'entity':'Entity not found.'})

@api_view(['GET','PUT'])
def regulatory_config(request, slug):
    entity=_entity(slug)
    config,_=RegulatoryConfiguration.objects.get_or_create(entity=entity,defaults={'reporting_currency':entity.base_currency,'fx_to_reporting':{entity.base_currency:'1'}})
    if request.method=='PUT':
        serializer=RegulatoryConfigurationSerializer(config,data=request.data,partial=True); serializer.is_valid(raise_exception=True); serializer.save()
    return Response(RegulatoryConfigurationSerializer(config).data)

@api_view(['GET'])
def regulatory_mappings(request, slug):
    return Response(RegulatoryMappingSerializer(RegulatoryMapping.objects.filter(entity=_entity(slug)),many=True).data)

@api_view(['PATCH'])
def regulatory_mapping_detail(request, slug, mapping_id):
    mapping=RegulatoryMapping.objects.filter(entity=_entity(slug),pk=mapping_id).first()
    if not mapping: raise ValidationError({'mapping_id':'Regulatory mapping not found.'})
    serializer=RegulatoryMappingSerializer(mapping,data=request.data,partial=True); serializer.is_valid(raise_exception=True); serializer.save()
    return Response(serializer.data)

@api_view(['POST','GET'])
def regulatory_calculations(request, slug):
    entity=_entity(slug)
    if request.method=='GET':
        qs=RegulatoryCalculation.objects.filter(entity=entity)[:25]
        return Response(RegulatoryCalculationSerializer(qs,many=True).data)
    config,_=RegulatoryConfiguration.objects.get_or_create(entity=entity,defaults={'reporting_currency':entity.base_currency,'fx_to_reporting':{entity.base_currency:'1'}})
    as_of=entity.configuration.as_of_date
    contracts=list(entity.portfolio_contracts.all())
    mappings=list(entity.regulatory_mappings.all())
    result=calculate_regulatory(entity,contracts,mappings,config,as_of)
    with transaction.atomic():
        calc=RegulatoryCalculation.objects.create(owner=request.user,entity=entity,as_of_date=as_of,engine_version=REGULATORY_ENGINE_VERSION,
            methodology_version=config.methodology_version,lcr_result=result['lcr'],nsfr_result=result['nsfr'],controls=result['controls'],warnings=result['warnings'])
        RegulatoryContribution.objects.bulk_create([RegulatoryContribution(calculation=calc,**row) for row in result['contributions']],batch_size=1000)
    return Response(RegulatoryCalculationSerializer(calc).data,status=201)

@api_view(['GET'])
def regulatory_calculation_detail(request, slug, calculation_id):
    calc=RegulatoryCalculation.objects.filter(entity=_entity(slug),pk=calculation_id).first()
    if not calc: raise ValidationError({'calculation_id':'Regulatory calculation not found.'})
    return Response(RegulatoryCalculationSerializer(calc).data)

@api_view(['GET'])
def regulatory_contributions(request, slug, calculation_id):
    calc=RegulatoryCalculation.objects.filter(entity=_entity(slug),pk=calculation_id).first()
    if not calc: raise ValidationError({'calculation_id':'Regulatory calculation not found.'})
    qs=calc.contributions.all()
    metric=request.query_params.get('metric'); line=request.query_params.get('line')
    if metric: qs=qs.filter(metric=metric.upper())
    if line: qs=qs.filter(category_code=line)
    return Response(RegulatoryContributionSerializer(qs[:2000],many=True).data)

@api_view(['GET'])
def regulatory_export(request, slug, calculation_id, metric):
    calc=RegulatoryCalculation.objects.filter(entity=_entity(slug),pk=calculation_id).first()
    if not calc: raise ValidationError({'calculation_id':'Regulatory calculation not found.'})
    metric=metric.upper()
    if metric not in ('LCR','NSFR'): raise ValidationError({'metric':'Choose LCR or NSFR.'})
    result=calc.lcr_result if metric=='LCR' else calc.nsfr_result
    response=HttpResponse(content_type='text/csv')
    response['Content-Disposition']=f'attachment; filename="{metric.lower()}-{slug}-{calc.as_of_date}.csv"'
    writer=csv.writer(response)
    writer.writerow([metric,'Report date',calc.as_of_date,'Reporting currency',calc.controls.get('reporting_currency','')])
    writer.writerow([]); writer.writerow(['Section','Code','Line item','Source balance','Weighted amount'])
    for row in result.get('lines',[]): writer.writerow([row.get('section'),row.get('code'),row.get('label'),row.get('source_balance'),row.get('weighted_amount')])
    writer.writerow([])
    if metric=='LCR':
        for label,key in [('HQLA','hqla'),('Gross outflows','gross_outflows'),('Gross inflows','gross_inflows'),('Eligible inflows','eligible_inflows'),('Net cash outflows','net_cash_outflows'),('LCR %','ratio')]: writer.writerow([label,result.get(key)])
    else:
        for label,key in [('ASF','asf'),('RSF','rsf'),('NSFR %','ratio')]: writer.writerow([label,result.get(key)])
    return response

@api_view(['GET'])
def regulatory_workbook(request, slug, calculation_id):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    calc=RegulatoryCalculation.objects.filter(entity=_entity(slug),pk=calculation_id).first()
    if not calc: raise ValidationError({'calculation_id':'Regulatory calculation not found.'})
    wb=Workbook(); wb.remove(wb.active)
    header_fill=PatternFill('solid',fgColor='D9EAF2'); title_fill=PatternFill('solid',fgColor='18364A')
    def report_sheet(name,result):
        ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
        ws['A1']=f'{name} regulatory report'; ws['A1'].font=Font(bold=True,color='FFFFFF',size=14); ws['A1'].fill=title_fill
        ws.merge_cells('A1:E1'); ws['A2']='Entity';ws['B2']=calc.entity.name;ws['C2']='Report date';ws['D2']=calc.as_of_date.isoformat()
        headers=['Section','Code','Line item','Source balance','Weighted amount']
        for col,h in enumerate(headers,1): c=ws.cell(4,col,h);c.font=Font(bold=True);c.fill=header_fill
        row=5
        for item in result.get('lines',[]):
            vals=[item.get('section'),item.get('code'),item.get('label'),float(item.get('source_balance',0)),float(item.get('weighted_amount',0))]
            for col,v in enumerate(vals,1): ws.cell(row,col,v)
            row+=1
        row+=1
        summary = ([('HQLA','hqla'),('Gross outflows','gross_outflows'),('Gross inflows','gross_inflows'),('Eligible inflows','eligible_inflows'),('Net cash outflows','net_cash_outflows'),('LCR %','ratio')]
                   if name=='LCR' else [('Available stable funding','asf'),('Required stable funding','rsf'),('NSFR %','ratio')])
        for label,key in summary:
            ws.cell(row,3,label).font=Font(bold=True); raw=result.get(key); ws.cell(row,5,None if raw is None else float(raw)); row+=1
        ws.column_dimensions['A'].width=16;ws.column_dimensions['B'].width=24;ws.column_dimensions['C'].width=48;ws.column_dimensions['D'].width=20;ws.column_dimensions['E'].width=20
        for r in ws.iter_rows(min_row=5,min_col=4,max_col=5):
            for c in r:c.number_format='#,##0.000'
    report_sheet('LCR',calc.lcr_result);report_sheet('NSFR',calc.nsfr_result)
    ws=wb.create_sheet('Controls & Warnings');ws.sheet_view.showGridLines=False
    ws.append(['Control','Value']);
    for k,v in calc.controls.items(): ws.append([k,str(v)])
    ws.append([]);ws.append(['Contract','Type','Warning'])
    for w in calc.warnings: ws.append([w.get('contract_id'),w.get('type'),w.get('message')])
    ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=24;ws.column_dimensions['C'].width=80
    audit=wb.create_sheet('Audit Contributions');audit.sheet_view.showGridLines=False
    headers=['Metric','Contract ID','Product','Source CCY','Reporting balance','Line code','Line item','Factor','Weighted amount','Maturity band']
    audit.append(headers)
    for cell in audit[1]: cell.font=Font(bold=True);cell.fill=header_fill
    for c in calc.contributions.all().iterator():
        audit.append([c.metric,c.contract_id,c.product,c.currency,float(c.source_balance),c.category_code,c.category_label,float(c.factor),float(c.weighted_amount),c.maturity_band])
    widths=[10,22,20,12,20,24,48,12,20,16]
    for i,w in enumerate(widths,1): audit.column_dimensions[chr(64+i)].width=w
    out=BytesIO();wb.save(out);out.seek(0)
    response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=f'attachment; filename="lcr-nsfr-{slug}-{calc.as_of_date}.xlsx"'
    return response
