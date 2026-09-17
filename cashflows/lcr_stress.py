"""Source-driven implementation of the legacy LCR Stress Testing rule library."""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from .regulatory import lcr_report_from_positions

D=Decimal
DEFAULTS=[
 {'id':'deposit_withdrawal','name':'Deposit Withdrawal','description':'This shock simulates a situation where a percentage of the bank’s deposits is withdrawn, increasing stressed cash outflows.','enabled':True,'system_default':True,'protected':True,'levels':[{'id':'moderate','label':'10%'},{'id':'medium','label':'20%'},{'id':'severe','label':'30%'}],'rules':[{'name':'Stable retail and SME deposits','target':'stable_deposits','operation':'rate_floor','values':{'moderate':.05,'medium':.10,'severe':.20}},{'name':'Less-stable retail and SME deposits','target':'less_stable_deposits','operation':'rate_floor','values':{'moderate':.10,'medium':.20,'severe':.30}}]},
 {'id':'unused_facilities','name':'Unused Facilities Withdrawal','description':'This scenario assumes that a percentage of unused committed credit and liquidity facilities is drawn.','enabled':True,'system_default':True,'protected':True,'levels':[{'id':'moderate','label':'20%'},{'id':'medium','label':'40%'},{'id':'severe','label':'100%'}],'rules':[{'name':'Committed credit and liquidity facilities','target':'unused_facilities','operation':'rate_floor','values':{'moderate':.20,'medium':.40,'severe':1}}]},
 {'id':'top_depositors','name':'Top Depositors Withdrawal','description':'This shock considers the effect of withdrawals by the bank’s largest depositors using aggregate withdrawal inputs.','enabled':True,'system_default':True,'protected':True,'levels':[{'id':'moderate','label':'Top 1'},{'id':'medium','label':'Top 3'},{'id':'severe','label':'Top 5'}],'rules':[{'name':'Aggregate top-depositor withdrawal','target':'manual_top_depositors','operation':'additional_outflow','value_source':'manual_top_depositors','values':{'moderate':0,'medium':0,'severe':0}}]},
 {'id':'hqla_haircut','name':'HQLA Haircut','description':'A percentage haircut is applied to the value of high-quality liquid assets.','enabled':True,'system_default':True,'protected':True,'levels':[{'id':'moderate','label':'10%'},{'id':'medium','label':'20%'},{'id':'severe','label':'40%'}],'rules':[{'name':'Reduce HQLA component amounts','target':'hqla_components','operation':'amount_haircut','values':{'moderate':.10,'medium':.20,'severe':.40}}]},
 {'id':'credit_default','name':'Credit Default','description':'This shock applies a percentage credit-default rate to expected cash inflows from fully performing loans.','enabled':True,'system_default':True,'protected':True,'levels':[{'id':'moderate','label':'10%'},{'id':'medium','label':'15%'},{'id':'severe','label':'20%'}],'rules':[{'name':'Reduce performing-loan inflows','target':'credit_inflows','operation':'amount_haircut','values':{'moderate':.10,'medium':.15,'severe':.20}}]},
]
DEFAULT_TOP={'moderate':180000,'medium':420000,'severe':680000}
OPERATIONS={'rate_floor','rate_cap','rate_set','rate_add','rate_subtract','rate_increase_pct','rate_decrease_pct','amount_haircut','amount_decrease_pct','amount_increase_pct','amount_multiplier','amount_add','amount_subtract','amount_set','set_zero','additional_outflow','additional_inflow'}

def default_configuration(): return {'version':'2.0','inflow_cap':.75,'lcr_limit':1,'lcr_tolerance':1.2,'scenarios':deepcopy(DEFAULTS)}
def normalise(config):
    value=deepcopy(config or default_configuration());value.setdefault('version','2.0');value.setdefault('inflow_cap',.75);value.setdefault('lcr_limit',1);value.setdefault('lcr_tolerance',1.2);scenarios=value.setdefault('scenarios',[]);found={x.get('id') for x in scenarios}
    for scenario in DEFAULTS:
        if scenario['id'] not in found: scenarios.append(deepcopy(scenario))
    return value
def groups(position):
    category=position.get('lcr_category','');key=position.get('key','')
    output={'all_input_elements'}
    if category.startswith('level'): output|={'hqla_components'}
    if category=='retail_stable': output|={'stable_deposits','outflows'}
    if category=='retail_less_stable': output|={'less_stable_deposits','outflows'}
    if category in ('wholesale_operational','wholesale_non_operational','secured_funding','other_outflows'): output|={'outflows'}
    if category=='other_outflows' or key in ('unused_facilities','contingent'): output|={'unused_facilities'}
    if category.startswith('inflow'): output|={'inflows'}
    if category=='inflow_customer' or key in ('customer_inflows','credit_inflows'): output|={'credit_inflows'}
    return output
def value(rule,level): return D(str(rule.get('values',{}).get(level,0)))
def apply_operation(position,operation,amount):
    balance=D(str(position.get('balance',0)));factor=D(str(position.get('lcr_factor',0)))
    if operation=='rate_floor': factor=max(factor,amount)
    elif operation=='rate_cap': factor=min(factor,amount)
    elif operation=='rate_set': factor=max(D(0),amount)
    elif operation=='rate_add': factor=max(D(0),factor+amount)
    elif operation=='rate_subtract': factor=max(D(0),factor-amount)
    elif operation=='rate_increase_pct': factor=max(D(0),factor*(1+amount))
    elif operation=='rate_decrease_pct': factor=max(D(0),factor*(1-amount))
    elif operation in ('amount_haircut','amount_decrease_pct'): balance=max(D(0),balance*(1-amount))
    elif operation=='amount_increase_pct': balance=max(D(0),balance*(1+amount))
    elif operation=='amount_multiplier': balance=max(D(0),balance*amount)
    elif operation=='amount_add': balance=max(D(0),balance+amount)
    elif operation=='amount_subtract': balance=max(D(0),balance-amount)
    elif operation=='amount_set': balance=max(D(0),amount)
    elif operation=='set_zero': balance=D(0)
    else: raise ValueError(f'Unsupported operation: {operation}')
    position['balance']=str(balance);position['lcr_factor']=str(factor)
def calculate(entity,as_of,source,configuration,top_amounts):
    config=normalise(configuration);baseline=lcr_report_from_positions(source,entity,as_of_date=as_of,basis='Source-driven LCR stress testing baseline.')
    baseline_lcr=D(baseline['lcr']);results=[];audit=[]
    for scenario in (x for x in config['scenarios'] if x.get('enabled',True)):
        for level in scenario.get('levels',[]):
            positions=deepcopy(source);additional_outflow=D(0);additional_inflow=D(0)
            for rule in scenario.get('rules',[]):
                operation=rule.get('operation');amount=value(rule,level['id'])
                if operation in ('additional_outflow','additional_inflow'):
                    applied=D(str(top_amounts.get(level['id'],0))) if rule.get('value_source')=='manual_top_depositors' or rule.get('target')=='manual_top_depositors' else amount
                    if operation=='additional_outflow': additional_outflow+=applied
                    else: additional_inflow+=applied
                    audit.append({'scenario':scenario['name'],'severity':level['label'],'rule':rule.get('name',''),'operation':operation,'target':rule.get('target',''),'element':'Independent cash-flow adjustment','baseline_amount':'0','stressed_amount':str(applied),'weighted_delta':str(applied if operation=='additional_outflow' else -applied)})
                    continue
                for position in positions:
                    selected=rule.get('target') in groups(position) or position.get('key') in rule.get('element_ids',[]) or position.get('lcr_category') in rule.get('element_ids',[])
                    if not selected: continue
                    before=D(str(position.get('balance',0)))*D(str(position.get('lcr_factor',0)));old_balance=position.get('balance',0);old_factor=position.get('lcr_factor',0)
                    apply_operation(position,operation,amount);after=D(str(position['balance']))*D(str(position['lcr_factor']))
                    if before!=after: audit.append({'scenario':scenario['name'],'severity':level['label'],'rule':rule.get('name',''),'operation':operation,'target':rule.get('target',''),'element':position.get('label') or position.get('key'),'baseline_amount':str(old_balance),'baseline_rate':str(old_factor),'stressed_amount':position['balance'],'stressed_rate':position['lcr_factor'],'weighted_delta':str(after-before)})
            if additional_outflow: positions.append({'key':'stress_additional_outflow','label':'Independent additional cash outflow','balance':str(additional_outflow),'lcr_category':'other_outflows','lcr_factor':'1'})
            if additional_inflow: positions.append({'key':'stress_additional_inflow','label':'Independent additional cash inflow','balance':str(additional_inflow),'lcr_category':'inflow_customer','lcr_factor':'1'})
            report=lcr_report_from_positions(positions,entity,as_of_date=as_of,basis='Deterministic source-driven LCR stress calculation.')
            lcr=D(report['lcr']);limit=D(str(config['lcr_limit']));tolerance=D(str(config['lcr_tolerance']));status='Below Limit' if lcr<limit else 'Tolerance' if lcr<tolerance else 'Within Appetite'
            results.append({'scenario_id':scenario['id'],'scenario':scenario['name'],'description':scenario.get('description',''),'type':'System Default' if scenario.get('system_default') else 'Custom','severity_id':level['id'],'severity':level['label'],'hqla':report['hqla'],'weighted_outflows':report['total_outflows'],'gross_weighted_inflows':report['total_inflows'],'inflow_cap':str(D(report['total_outflows'])*D(str(config['inflow_cap']))),'recognized_inflows':report['eligible_inflows'],'net_cash_outflow':report['net_cash_outflows'],'lcr':report['lcr'],'movement':str(lcr-baseline_lcr),'risk_status':status,'lcr_report':report})
    return {'baseline':{'hqla':baseline['hqla'],'weighted_outflows':baseline['total_outflows'],'gross_weighted_inflows':baseline['total_inflows'],'inflow_cap':str(D(baseline['total_outflows'])*D(str(config['inflow_cap']))),'recognized_inflows':baseline['eligible_inflows'],'net_cash_outflow':baseline['net_cash_outflows'],'lcr':baseline['lcr'],'risk_status':'Within Appetite' if baseline_lcr>=D(str(config['lcr_tolerance'])) else 'Tolerance' if baseline_lcr>=D(str(config['lcr_limit'])) else 'Below Limit'},'results':results,'audit':audit,'limit':str(config['lcr_limit']),'tolerance':str(config['lcr_tolerance']),'inflow_cap':str(config['inflow_cap'])}
def xlsx(run, entity_name='Jordan', as_of_date=None):
    """Export the familiar bank-format stress overview, plus full results and audit sheets."""
    results=run['results'];wb=Workbook();ws=wb.active;ws.title='Summary'
    navy='183B57';red='D90000';peach='F4B183';green='70AD47';line='AAB4C3';thin=Side(style='thin',color=line)
    border=Border(left=thin,right=thin,top=thin,bottom=thin)
    centre=Alignment(horizontal='center',vertical='center');left=Alignment(vertical='center',wrap_text=True)
    title_date=as_of_date.strftime('%B %Y') if as_of_date else 'Current'
    ws.merge_cells('A1:B1');ws.merge_cells('C1:D1');ws['A1']=f'LCR ST {title_date}';ws['C1']=entity_name
    for cell in ('A1','C1'):
        ws[cell].fill=PatternFill('solid',fgColor=navy);ws[cell].font=Font(bold=True,color='FFFFFF',size=12);ws[cell].alignment=centre
    thresholds=[('Limit',results['limit'],red),('Tolerance',results['tolerance'],peach),('Risk Appetite',f">{results['tolerance']}",green)]
    for row,(label,value,color) in enumerate(thresholds,2):
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2);ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=4)
        ws.cell(row,1,label);ws.cell(row,3,value)
        for col in range(1,5):
            cell=ws.cell(row,col);cell.fill=PatternFill('solid',fgColor=color);cell.font=Font(bold=True,color='FFFFFF' if color==red else '10233E');cell.alignment=centre
    ws['A5']='Baseline';ws['B5']='';ws['C5']=float(results['baseline']['lcr']);ws['D5']=''
    ws.merge_cells('A5:B5')
    for cell in ws[5]: cell.border=border;cell.font=Font(bold=True);cell.alignment=centre
    ws['C5'].number_format='0.0%'
    headers=['Scenario','Severity','LCR ST','Movement']
    for col,header in enumerate(headers,1):
        cell=ws.cell(6,col,header);cell.fill=PatternFill('solid',fgColor=navy);cell.font=Font(bold=True,color='FFFFFF');cell.alignment=centre;cell.border=border
    row_index=7;groups=[]
    for item in results['results']:
        if not groups or groups[-1][0]!=item['scenario']:
            groups.append((item['scenario'],[]))
        groups[-1][1].append(item)
    for scenario,items in groups:
        start=row_index
        for item in items:
            ws.cell(row_index,2,item['severity']);ws.cell(row_index,3,float(item['lcr']));ws.cell(row_index,4,float(item['movement']))
            for col in range(1,5): ws.cell(row_index,col).border=border;ws.cell(row_index,col).alignment=centre
            ws.cell(row_index,3).number_format='0.0%';ws.cell(row_index,4).number_format='0.0%;-0.0%'
            row_index+=1
        if len(items)>1: ws.merge_cells(start_row=start,start_column=1,end_row=row_index-1,end_column=1)
        ws.cell(start,1,scenario);ws.cell(start,1).font=Font(bold=True);ws.cell(start,1).alignment=left;ws.cell(start,1).border=border
    row_index+=2
    for col,header in enumerate(['Scenario','Description'],1):
        cell=ws.cell(row_index,col,header);cell.fill=PatternFill('solid',fgColor=navy);cell.font=Font(bold=True,color='FFFFFF');cell.alignment=centre;cell.border=border
    ws.merge_cells(start_row=row_index,start_column=1,end_row=row_index,end_column=2);ws.merge_cells(start_row=row_index,start_column=3,end_row=row_index,end_column=4)
    row_index+=1
    ws.merge_cells(start_row=row_index,start_column=1,end_row=row_index,end_column=2);ws.merge_cells(start_row=row_index,start_column=3,end_row=row_index,end_column=4)
    ws.cell(row_index,1,'Baseline');ws.cell(row_index,3,'This scenario represents normal operating conditions without any stress.')
    for col in range(1,5): ws.cell(row_index,col).border=border;ws.cell(row_index,col).alignment=left
    row_index+=1
    for scenario,items in groups:
        ws.merge_cells(start_row=row_index,start_column=1,end_row=row_index,end_column=2);ws.merge_cells(start_row=row_index,start_column=3,end_row=row_index,end_column=4)
        ws.cell(row_index,1,scenario);ws.cell(row_index,3,items[0]['description'])
        for col in range(1,5): ws.cell(row_index,col).border=border;ws.cell(row_index,col).alignment=left
        ws.row_dimensions[row_index].height=42;row_index+=1
    for column,width in {'A':28,'B':16,'C':18,'D':18}.items(): ws.column_dimensions[column].width=width
    rs=wb.create_sheet('Results');rs.append(['Scenario','Severity','HQLA','Weighted Outflows','Gross Inflows','75% Inflow Cap','Recognized Inflows','Net Cash Outflow','LCR','Movement','Risk Status'])
    for item in results['results']: rs.append([item['scenario'],item['severity'],float(item['hqla']),float(item['weighted_outflows']),float(item['gross_weighted_inflows']),float(item['inflow_cap']),float(item['recognized_inflows']),float(item['net_cash_outflow']),float(item['lcr']),float(item['movement']),item['risk_status']])
    audit=wb.create_sheet('Rule Impact Audit');audit.append(['Scenario','Severity','Rule','Operation','Target','Element','Baseline Amount','Stressed Amount','Weighted Delta'])
    for item in results['audit']: audit.append([item.get(key,'') for key in ('scenario','severity','rule','operation','target','element','baseline_amount','stressed_amount','weighted_delta')])
    for sheet in (rs,audit):
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor=navy);cell.alignment=centre
        for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=min(42,max(12,max(len(str(cell.value or '')) for cell in column)+2))
    output=BytesIO();wb.save(output);return output.getvalue()
